import os
from typing import Dict, List, Any


class DocxExtractor:
    """Extractor for DOCX files using python-docx with table and paragraph parsing."""

    @staticmethod
    def extract(file_path: str) -> Dict[str, Any]:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"DOCX file not found: {file_path}")

        try:
            import docx

            doc = docx.Document(file_path)
            content_blocks: List[str] = []

            # Extract paragraphs
            for p in doc.paragraphs:
                text = p.text.strip()
                if text:
                    content_blocks.append(text)

            # Extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        content_blocks.append(" | ".join(row_cells))

            full_text = "\n\n".join(content_blocks)

            # Estimate pagination (~3000 chars per page)
            page_size = 3000
            pages = []
            if not full_text.strip():
                pages = [{"page_number": 1, "text": ""}]
            else:
                chunks = [full_text[i:i + page_size] for i in range(0, len(full_text), page_size)]
                for idx, chunk in enumerate(chunks):
                    pages.append({"page_number": idx + 1, "text": chunk})

            return {
                "full_text": full_text,
                "page_count": len(pages),
                "pages": pages,
                "extractor": "python-docx",
            }
        except Exception as e:
            raise RuntimeError(f"Error extracting text from DOCX file: {str(e)}")
